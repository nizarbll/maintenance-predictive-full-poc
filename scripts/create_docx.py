import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_beautiful_docx():
    doc = Document()

    # --- Setup Global Styles ---
    # Normal Text
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(12.5)  # Increased from 11 to 12.5 for better readability
    
    # Custom Title Style
    title_style = doc.styles.add_style('PremiumTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_font = title_style.font
    title_font.name = 'Arial'
    title_font.size = Pt(26)
    title_font.bold = True
    title_font.color.rgb = RGBColor(30, 60, 114) # Deep Blue
    
    # Custom Subtitle Style
    sub_style = doc.styles.add_style('PremiumSubtitle', WD_STYLE_TYPE.PARAGRAPH)
    sub_font = sub_style.font
    sub_font.name = 'Arial'
    sub_font.size = Pt(16)
    sub_font.italic = True
    sub_font.color.rgb = RGBColor(42, 82, 152)

    # Heading 1
    h1_style = doc.styles['Heading 1']
    h1_font = h1_style.font
    h1_font.name = 'Arial'
    h1_font.size = Pt(18)
    h1_font.color.rgb = RGBColor(30, 60, 114)
    h1_font.bold = True

    # Heading 2
    h2_style = doc.styles['Heading 2']
    h2_font = h2_style.font
    h2_font.name = 'Arial'
    h2_font.size = Pt(14)
    h2_font.color.rgb = RGBColor(42, 82, 152)
    h2_font.bold = True

    def add_justified_paragraph(text, bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.bold = bold
        return p

    # --- Page de Garde ---
    doc.add_paragraph('\n\n\n')
    t = doc.add_paragraph('Rapport Technique', style='PremiumTitle')
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    st = doc.add_paragraph('Projet de Machine Learning Industriel', style='PremiumSubtitle')
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('\n\n')
    
    # Big Highlight box for TwinSense AI
    twin_title = doc.add_paragraph('TwinSense AI', style='PremiumTitle')
    twin_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    twin_title.runs[0].font.size = Pt(36)
    
    twin_sub = doc.add_paragraph('Architecture DataOps & MLOps pour la Maintenance Prédictive (NASA CMAPSS)', style='PremiumSubtitle')
    twin_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('\n\n\n\n')
    
    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = team.add_run('Équipe Projet :\nNizar Baloubali\nManal Es-sobhy\nAnas El Maslouti\nNouamane Chaibi\nAmine Ait Hamma\n\n')
    r.bold = True
    r.font.size = Pt(12)
    
    encad = team.add_run('Encadré par : Mr. El Hajji\nDate : Juin 2026')
    encad.font.color.rgb = RGBColor(42, 82, 152)
    encad.bold = True
    
    doc.add_page_break()

    # --- Résumé ---
    doc.add_heading('Résumé', level=1)
    add_justified_paragraph(
        "Ce rapport technique exhaustif détaille la conception, l'architecture, l'implémentation et les "
        "performances de la plateforme TwinSense AI. Face aux enjeux critiques de l'industrie aérospatiale, "
        "notamment l'optimisation des cycles de maintenance et la prévention des défaillances catastrophiques, "
        "ce projet vise à estimer précisément la Durée de Vie Résiduelle (RUL) de moteurs turbofan.\n\n"
        "S'appuyant sur les données historiques CMAPSS de la NASA, l'architecture déployée est une "
        "solution complète de bout en bout intégrant un pipeline DataOps (Apache Kafka), un Feature Store "
        "(Feast), un orchestrateur MLOps (MLflow) et un modèle d'apprentissage profond séquentiel (LSTM).\n\n"
        "Ce document décrit de manière approfondie les mécanismes de génération de données, le traitement "
        "en flux, la détection de dérive (Data Drift), le réentraînement continu, ainsi que l'intégration "
        "visuelle dans un tableau de bord industriel (Kennedy Space Center)."
    )
    doc.add_page_break()

    # --- 1. Introduction ---
    doc.add_heading('1. Introduction', level=1)
    add_justified_paragraph(
        "L'industrie aérospatiale est un secteur où l'erreur n'est pas tolérée. "
        "La maintenance réactive (attendre la panne) ou préventive (remplacements à intervalles fixes) ne "
        "suffisent plus. La maintenance prédictive, s'appuyant sur l'intelligence artificielle et l'IIoT, "
        "permet d'intervenir juste à temps, maximisant la disponibilité opérationnelle tout en réduisant les coûts."
    )
    doc.add_heading('1.1. Contexte du Jeu de Données NASA CMAPSS', level=2)
    add_justified_paragraph(
        "Le jeu de données CMAPSS modélise la dégradation de moteurs turbofan. "
        "Chaque moteur commence dans un état de santé normal et se dégrade au fil des cycles d'utilisation "
        "jusqu'à la défaillance totale (Run-to-Failure). Les données incluent 21 capteurs de télémétrie "
        "et 3 paramètres opérationnels."
    )
    
    # --- 2. Architecture ---
    doc.add_heading('2. Architecture et Solution Proposée', level=1)
    add_justified_paragraph(
        "L'architecture TwinSense AI repose sur une approche micro-services segmentée en plusieurs sous-systèmes interdépendants."
    )
    
    doc.add_heading('2.1. Pipeline DataOps : Kafka et Architecture Medaillon', level=2)
    add_justified_paragraph(
        "Un générateur Python produit des flux télémétriques JSON en temps réel poussés sur Apache Kafka. "
        "Le pipeline consomme ces flux et applique une architecture Medaillon : Bronze (Brut), Silver (Nettoyé), "
        "et Gold (Moyennes glissantes sur 5 cycles et écarts-types)."
    )
    
    doc.add_heading('2.2. Feature Store : Feast', level=2)
    add_justified_paragraph(
        "Feast garantit le Point-in-Time Correctness en sauvegardant les données avec horodatage, "
        "éliminant ainsi tout risque de fuite de données entre l'entraînement et l'inférence en production."
    )
    
    doc.add_heading('2.3. Réseau CNN-LSTM et Auto-Guérison MLflow', level=2)
    add_justified_paragraph(
        "Un modèle hybride (Conv1D + LSTM + Dense) analyse des séquences de 30 cycles pour prédire le RUL. "
        "MLflow surveille l'erreur en temps réel : si la MAE dérive de plus de 20%, un Challenger est entraîné "
        "et promu automatiquement sans intervention humaine."
    )

    doc.add_heading('2.4. Intégration Interface Homme-Machine', level=2)
    add_justified_paragraph(
        "Le tableau de bord développé (TwinSense AI) offre un rendu visuel impressionnant avec une carte "
        "interactive du Kennedy Space Center et des visualisations de dégradation des capteurs via Chart.js."
    )
    
    # Insert Screenshot if exists
    img_path = os.path.join(os.path.dirname(__file__), '..', 'presentation', 'frontend_screenshot.png')
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(img_path, width=Inches(6.0))
        p_cap = doc.add_paragraph('Figure 1 : Aperçu des prédictions (simulées) sur le tableau de bord', style='Caption')
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # --- 3. Résultats & Limites ---
    doc.add_heading('3. Résultats et Limites Techniques', level=1)
    
    doc.add_heading('3.1. Performances', level=2)
    add_justified_paragraph(
        "Le modèle Baseline converge vers une MAE d'environ 12 cycles. De plus, le système a réussi notre "
        "Stress Test : lors de la simulation d'une usure accélérée anormale, MLflow a capté la dérive et a "
        "auto-régénéré un modèle Challenger en production."
    )
    
    doc.add_heading('3.2. Le Paradoxe du RUL en Streaming', level=2)
    add_justified_paragraph(
        "Afin de fournir des labels 'vrais' pour l'entraînement MLflow en streaming, le RUL a été calculé de "
        "manière déterministe par le producteur. Dans une véritable implémentation industrielle, on ignorerait "
        "le RUL futur. Le système devrait donc scinder l'entraînement Batch asynchrone (sur moteurs cassés) "
        "et l'inférence pur Streaming."
    )

    doc.add_heading('4. Conclusion', level=1)
    add_justified_paragraph(
        "Le projet TwinSense AI illustre la synergie indispensable entre le Data Engineering et le MLOps. "
        "Couplée à un tableau de bord industriel réactif, l'Intelligence Artificielle quitte les carnets de "
        "recherche pour devenir un outil d'aide à la décision. TwinSense AI propose un paradigme où la "
        "maintenance devient proactive, prescriptive et autonome."
    )

    doc.add_page_break()

    # --- 5. Références ---
    doc.add_heading('5. Références', level=1)
    
    def add_reference(text):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(text).italic = True

    add_reference("NASA CMAPSS Dataset : A. Saxena, K. Goebel, D. Simon, and N. Eklund, 'Damage Propagation Modeling for Aircraft Engine Run-to-Failure Simulation', PHM08, Denver CO, Oct 2008.")
    add_reference("Deep Learning Sérentiel : Hochreiter, S., & Schmidhuber, J. (1997). Long Short-Term Memory. Neural Computation, 9(8), 1735–1780.")
    add_reference("DataOps - Apache Kafka : Documentation officielle d'Apache Kafka. (kafka.apache.org)")
    add_reference("Feature Store : Feast (Feature Store) Concepts et Architecture. (docs.feast.dev)")
    add_reference("MLOps : Documentation Officielle MLflow - A Platform for the Machine Learning Lifecycle. (mlflow.org)")

    output_path = os.path.join(os.path.dirname(__file__), '..', 'presentation', 'ML-Indus-rapport.docx')
    doc.save(output_path)
    print(f"Beautiful Document saved as {output_path}")

if __name__ == "__main__":
    create_beautiful_docx()
