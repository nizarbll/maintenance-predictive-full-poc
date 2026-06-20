import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_premium_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    def add_styled_slide():
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Clean solid background (Apple/Stripe minimalist style)
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(248, 249, 250) # Very light cool gray
        
        return slide

    def add_title_box(slide, text, top=0.6, size=32):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.733), Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = 'Segoe UI' # Modern font, looks great on Windows
        p.font.bold = True
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(10, 37, 64) # Deep premium navy blue
        
        # Add a subtle accent line under the title
        line = slide.shapes.add_connector(
            1, Inches(0.8), Inches(top + 0.8), Inches(3.0), Inches(top + 0.8)
        )
        line.line.color.rgb = RGBColor(0, 85, 255) # Bright accent blue
        line.line.width = Pt(2)
        
        return txBox

    def add_text_box(slide, text, top=1.8, size=18, bold=False):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.733), Inches(5))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        paragraphs = text.split('\n')
        for i, line_text in enumerate(paragraphs):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = line_text
            p.font.name = 'Segoe UI'
            p.font.size = Pt(size)
            p.font.bold = bold
            # Use a softer dark grey for reading
            p.font.color.rgb = RGBColor(60, 66, 87)
            # Line spacing
            p.space_before = Pt(8)
            p.space_after = Pt(4)
        return tf

    # SLIDE 1 : Titre
    s1 = add_styled_slide()
    
    # Big Title
    tb_title = s1.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    p = tb_title.text_frame.paragraphs[0]
    p.text = "TwinSense AI"
    p.font.name = 'Segoe UI'
    p.font.bold = True
    p.font.size = Pt(64)
    p.font.color.rgb = RGBColor(0, 85, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    tb_sub = s1.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(12.333), Inches(1))
    p_sub = tb_sub.text_frame.paragraphs[0]
    p_sub.text = "ARCHITECTURE DATAOPS & MLOPS POUR LA MAINTENANCE PRÉDICTIVE"
    p_sub.font.name = 'Segoe UI'
    p_sub.font.bold = True
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(10, 37, 64)
    p_sub.alignment = PP_ALIGN.CENTER
    
    # Team
    tb_team = s1.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.333), Inches(1))
    p_team = tb_team.text_frame.paragraphs[0]
    p_team.text = "Nizar Baloubali • Manal Es-sobhy • Anas El Maslouti • Nouamane Chaibi • Amine Ait Hamma\nEncadré par : Mr. El Hajji"
    p_team.font.name = 'Segoe UI'
    p_team.font.size = Pt(14)
    p_team.font.color.rgb = RGBColor(100, 110, 130)
    p_team.alignment = PP_ALIGN.CENTER

    # SLIDE 2 : Enjeu Industriel
    s2 = add_styled_slide()
    add_title_box(s2, "1. L'Enjeu Industriel de l'Aérospatiale")
    add_text_box(s2, "Le secteur aérospatial ne tolère aucune marge d'erreur.\nLes défaillances de moteurs en plein vol entraînent des conséquences catastrophiques.\n\n• Maintenance Réactive (Run-to-Failure) : Inenvisageable sur des turbofans critiques.\n• Maintenance Préventive : Remplacement à intervalles fixes gaspille des pièces saines.\n\nNotre Solution : La Maintenance Prédictive.\nPrédire la Durée de Vie Résiduelle (RUL) grâce au Machine Learning.")

    # SLIDE 3 : Dataset CMAPSS
    s3 = add_styled_slide()
    add_title_box(s3, "2. Fondation des Données : NASA CMAPSS")
    add_text_box(s3, "Pour simuler fidèlement ce comportement, nous avons exploité le Commercial Modular Aero-Propulsion System Simulation de la NASA.\n\n• Variables Opérationnelles : Altitude, Mach, Angle de la manette des gaz.\n• Télémétrie : 21 capteurs de haute précision (Températures, pressions, régimes moteurs).\n• Objectif MLOps : Observer la dégradation thermodynamique latente, invisible à l'œil nu, et la modéliser mathématiquement.")

    # SLIDE 4 : Génération de Flux
    s4 = add_styled_slide()
    add_title_box(s4, "3. DataOps : Le Générateur NASA")
    add_text_box(s4, "Dans le monde industriel, les données ne sont pas de simples fichiers plats. Ce sont des flux temporels.\n\n• scripts/generate_data.py : Un générateur stochastique qui simule des centaines de trajectoires de turbofans.\n• Bruit Gaussien : Des perturbations aléatoires sont injectées dans les signaux.\n• L'Output : Des objets JSON envoyés à très haute fréquence vers notre infrastructure.")

    # SLIDE 5 : Apache Kafka
    s5 = add_styled_slide()
    add_title_box(s5, "4. DataOps : L'Ingestion Kafka")
    add_text_box(s5, "Pour absorber le choc des capteurs envoyant des données à la milliseconde, l'architecture repose sur Apache Kafka.\n\n• Producer : Joue le rôle du moteur de l'avion, publiant ses événements en direct.\n• Consumer Pipeline : S'abonne au topic et traite la donnée par micro-batches asynchrones.\n• Résilience : Kafka garantit qu'aucune donnée n'est perdue, même si le serveur ML est indisponible.")

    # SLIDE 6 : Medaillon ETL
    s6 = add_styled_slide()
    add_title_box(s6, "5. DataOps : Architecture Medaillon")
    add_text_box(s6, "Les données ingérées par Kafka traversent un pipeline rigoureux :\n\n• Couche Bronze : Sauvegarde immédiate du flux brut (Traçabilité stricte).\n• Couche Silver : Nettoyage drastique (valeurs manquantes, anomalies schématiques).\n• Couche Gold : Feature Engineering avancé. Calcul de moyennes glissantes et d'écarts-types sur 5 cycles. C'est ce qui donne au réseau de neurones la notion de dynamique d'usure.")

    # SLIDE 7 : Feature Store Feast
    s7 = add_styled_slide()
    add_title_box(s7, "6. MLOps : Feast (Feature Store)")
    add_text_box(s7, "Un défi majeur en ML Industriel est le Training-Serving Skew.\n\n• Point-in-Time Correctness : Feast sauvegarde chaque caractéristique 'Gold' avec un horodatage précis (event_timestamp).\n• Prévention des Fuites : Cela évite le Data Leakage en garantissant que l'algorithme ne voit jamais de données 'du futur' lors de son entraînement historique.")

    # SLIDE 8 : Le Réseau LSTM
    s8 = add_styled_slide()
    add_title_box(s8, "7. Deep Learning : Réseau Hybride Conv1D-LSTM")
    add_text_box(s8, "Prédire le RUL nécessite de comprendre la séquence temporelle. Un réseau classique échouerait.\n\n• Couche Conv1D : Extrait les caractéristiques spatiales et lisse le bruit télémétrique brutal.\n• Couche LSTM : Long Short-Term Memory. Elle analyse l'historique des 30 derniers cycles.\n• Couches Denses : Avec 20% de Dropout pour projeter ces représentations vers une prédiction de RUL continue.")

    # SLIDE 9 : MLflow et Data Drift
    s9 = add_styled_slide()
    add_title_box(s9, "8. MLOps : Détection de Dérive (Data Drift)")
    add_text_box(s9, "L'IA en production devient obsolète si l'environnement physique change.\n\n• Évaluation Continue : MLflow surveille en permanence la Mean Absolute Error (MAE) du modèle en production.\n• Auto-Guérison : Si la MAE dépasse un seuil de tolérance (> 1.2x la MAE initiale), MLflow entraîne automatiquement un modèle Challenger.\n• Déploiement : Si le Challenger est meilleur, il est promu à la volée en Production.")

    # Helper function for image slides
    def add_image_slide(title, subtitle, img_name, img_width=8.0):
        s = add_styled_slide()
        add_title_box(s, title, top=0.5, size=28)
        
        # Subtitle
        tb = s.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(0.5))
        p = tb.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.name = 'Segoe UI'
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(100, 110, 130)
        
        img_path = os.path.join(os.path.dirname(__file__), '..', 'presentation', img_name)
        if os.path.exists(img_path):
            left_margin = (13.333 - img_width) / 2.0
            s.shapes.add_picture(img_path, Inches(left_margin), Inches(1.8), width=Inches(img_width))

    # SLIDES UI
    add_image_slide("8.1 Diagramme Architecture : DataOps & MLOps", "Flux d'ingestion Kafka, Feature Store et orchestration des Modèles.", "mlops_diagram.png", img_width=7.5)
    add_image_slide("9. Interface : TwinSense Dashboard", "Supervision Géospatiale Interactive au Kennedy Space Center.", "screenshot_dashboard.png")
    add_image_slide("10. Interface : Supervision de Flotte", "Registre dynamique des moteurs turbofans (ENG_001, ENG_002...).", "screenshot_machines.png")
    add_image_slide("11. Interface : Détails Télémétriques", "Croisement des capteurs CMAPSS et prédiction LSTM.", "frontend_screenshot.png")
    add_image_slide("12. Interface : Analyses Avancées", "Tableaux de bord historiques des pannes.", "screenshot_analytics.png")
    add_image_slide("13. Interface : Centre d'Alertes", "Notifications MLOps en temps réel et dispatch d'équipes.", "screenshot_alerts.png")

    # SLIDE 14 : Paradoxe du Streaming
    s14 = add_styled_slide()
    add_title_box(s14, "14. Discussion : Le Paradoxe du RUL")
    add_text_box(s14, "Le calcul du RUL a révélé un paradoxe industriel.\n\n• Le Biais de Simulation : Pour entraîner MLflow en temps réel, notre générateur a dû inclure le RUL futur dans le flux actuel.\n• La Réalité Industrielle : Il est mathématiquement impossible de connaître la panne future en pur temps réel.\n• Recommandation : L'architecture de production devrait isoler complètement l'Entraînement Batch (sur moteurs cassés) de l'Inférence Streaming.")

    # SLIDE 15 : Conclusion
    s15 = add_styled_slide()
    add_title_box(s15, "15. Conclusion & Perspectives")
    add_text_box(s15, "Le Machine Learning Industriel n'est pas qu'une question d'algorithmes mathématiques. C'est l'art de maintenir ces algorithmes fiables face au chaos physique.\n\n• Kubernetes : Pour gérer 500+ avions simultanément, le pipeline doit migrer vers des Pods K8s.\n• WebSockets : Connecter l'API Python TensorFlow au Frontend via WebSockets pour du véritable Push Data.\n\nTwinSense AI prouve qu'une télémétrie massive peut se transformer en décisions prescriptives autonomes.")

    out_path = os.path.join(os.path.dirname(__file__), '..', 'presentation', 'TwinSense_Presentation.pptx')
    prs.save(out_path)
    print(f"Premium Clean PPTX saved to {out_path}")

if __name__ == "__main__":
    create_premium_presentation()
